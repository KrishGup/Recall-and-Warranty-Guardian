"""Build the Guardian demo deck as PowerPoint from scripts/deck_content.py (the same content the video slides use).

    python scripts/build_deck.py                    # -> docs/deck/Guardian-Demo-Deck.pptx
    python scripts/build_deck.py --export           # also exports every slide to docs/deck/preview/NN.png through PowerPoint (Windows)

The deck is 16:9 at 13.333 x 7.5 in. The design (design/deck/README.md) is 1920x1080 px; PX() converts. Fonts: Roboto
Slab (titles, numbers), Lora (body), Habibi (eyebrows, labels), Consolas (identifiers). Install design/fonts/*.ttf for
the exact look; PowerPoint substitutes otherwise. Speaker notes carry the voice-over script."""
from __future__ import annotations

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from deck_content import AMBER, BLUE, CONNECTOR, GREEN, INK, LIVE, PAPER, RED, SLIDES, logo_svg  # noqa: E402

OUT_DIR = os.path.join(ROOT, "docs", "deck")
OUT = os.path.join(OUT_DIR, "Guardian-Demo-Deck.pptx")
SCREENS = os.path.join(ROOT, "design", "deck", "assets", "screens")
SW, SH = Inches(13.333), Inches(7.5)
SCALE = 13.333 / 1920  # inches per design pixel

HEAD = "Roboto Slab"
BODY = "Lora"
UI = "Habibi"
MONO = "Consolas"


def PX(v: float) -> Emu:
    return Inches(v * SCALE)


def FS(px: float) -> Pt:
    """Design px → points (13.333 in / 1920 px, 72 pt / in)."""
    return Pt(px * SCALE * 72)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


def rect(slide, x, y, w, h, fill: str, line: str | None = None, radius: float | None = None, line_w: float = 1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = min(0.5, radius / min(w, h)) if min(w, h) else 0.1
    shape.shadow.inherit = False
    return shape


def text(slide, x, y, w, h, runs, size: float, font: str = BODY, color: str = PAPER, bold: bool = False, italic: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing: float | None = None, caps: bool = False, spacing: float | None = None):
    """A text box. `runs` is a string or a list of (text, {overrides}) tuples; overrides: font, color, bold, italic, size."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    parts = [(runs, {})] if isinstance(runs, str) else runs
    for t, ov in parts:
        r = p.add_run()
        r.text = t.upper() if caps else t
        f = r.font
        f.name = ov.get("font", font)
        f.size = FS(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.italic = ov.get("italic", italic)
        f.color.rgb = rgb(ov.get("color", color))
        if spacing is not None:
            rPr = r._r.get_or_add_rPr()
            rPr.set("spc", str(int(spacing * 100)))
    return box


def paragraphs(slide, x, y, w, h, items: list[tuple[str, dict]], gap_pt: float = 8):
    """Several paragraphs in one box; each item is (text, {size, font, color, bold, italic, caps, line_spacing})."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for t, ov in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap_pt)
        if ov.get("line_spacing"):
            p.line_spacing = ov["line_spacing"]
        r = p.add_run()
        r.text = t.upper() if ov.get("caps") else t
        f = r.font
        f.name = ov.get("font", BODY)
        f.size = FS(ov.get("size", 28))
        f.bold = ov.get("bold", False)
        f.italic = ov.get("italic", False)
        f.color.rgb = rgb(ov.get("color", PAPER))
    return box


def hline(slide, x, y, w, color: str, weight: float = 1.0):
    ln = slide.shapes.add_connector(1, x, y, x + w, y)
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(weight)
    return ln


LOGO_PNG = os.path.join(OUT_DIR, "assets", "logo-nightwatch.png")


def render_logo() -> str:
    """The Night watch mark rendered once from the same markup as the web app (transparent PNG at 8x)."""
    if os.path.isfile(LOGO_PNG):
        return LOGO_PNG
    from playwright.sync_api import sync_playwright  # type: ignore

    os.makedirs(os.path.dirname(LOGO_PNG), exist_ok=True)
    size = 96
    html = f'<html><body style="margin:0;background:transparent"><div id="l" style="width:{size}px;height:{size}px">{logo_svg(size)}</div></body></html>'
    with sync_playwright() as p:
        b = p.chromium.launch()
        page = b.new_page(viewport={"width": size, "height": size}, device_scale_factor=8)
        page.set_content(html)
        page.locator("#l").screenshot(path=LOGO_PNG, omit_background=True)
        b.close()
    return LOGO_PNG


def logo(slide, x, y, size_px: float, dot_px: float, ring_px: float, word_px: float | None = None):
    """The Night watch mark (rendered image) with the wordmark beside it."""
    s = PX(size_px)
    slide.shapes.add_picture(render_logo(), x, y, s, s)
    if word_px:
        text(slide, x + s + PX(24), y, PX(400), s, "Guardian", word_px, font=HEAD, color=PAPER, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def background(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def notes(slide, s: dict):
    slide.notes_slide.notes_text_frame.text = s["notes"]


def add_image(slide, name: str, x, y, w=None, h=None, fallback: str | None = None):
    path = os.path.join(SCREENS, name)
    if not os.path.isfile(path) and fallback:
        path = os.path.join(SCREENS, fallback)
    if not os.path.isfile(path):
        return None
    if w is not None and h is not None:
        return slide.shapes.add_picture(path, x, y, w, h)
    if w is not None:
        return slide.shapes.add_picture(path, x, y, width=w)
    return slide.shapes.add_picture(path, x, y, height=h)


# ------------------------------------------------------------------------------------------------ slide builders
PADX, PADT, PADB = 110, 96, 90


def build_title(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    logo(sl, PX(PADX), PX(PADT), 72, 23, 5, word_px=44)
    text(sl, PX(PADX), PX(300), PX(1500), PX(300), s["h1"], 128, font=HEAD, bold=True, line_spacing=1.02)
    text(sl, PX(PADX), PX(600), PX(1300), PX(120), s["sub"], 40, font=BODY, line_spacing=1.3)
    text(sl, PX(PADX), PX(730), PX(1300), PX(60), s["tagline"], 34, font=BODY, italic=True, color="#B9C3BD")
    hline(sl, PX(PADX), PX(895), PX(1920 - 2 * PADX), "#2C3B33")
    text(sl, PX(PADX), PX(925), PX(1700), PX(40), s["footer_eyebrow"], 26, font=UI, color=AMBER, caps=True, spacing=2)
    text(sl, PX(PADX), PX(975), PX(1700), PX(40), s["footer_mono"], 26, font=MONO, color="#B9C3BD")
    notes(sl, s)


def build_problem(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    text(sl, PX(PADX), PX(PADT), PX(1200), PX(40), s["eyebrow"], 26, font=UI, color=AMBER, caps=True, spacing=2)
    text(sl, PX(PADX), PX(150), PX(1500), PX(200), s["h2"], 80, font=HEAD, bold=True, line_spacing=1.08)
    text(sl, PX(PADX), PX(380), PX(1240), PX(220), s["body"], 34, font=BODY, color="#B9C3BD", line_spacing=1.4)
    hline(sl, PX(PADX), PX(745), PX(1920 - 2 * PADX), "#2C3B33")
    colw = (1920 - 2 * PADX - 2 * 48) / 3
    for i, (n, c) in enumerate(s["stats"]):
        x = PX(PADX + i * (colw + 48))
        text(sl, x, PX(780), PX(colw), PX(110), n, 104, font=HEAD, bold=True, line_spacing=1.0)
        text(sl, x, PX(900), PX(colw), PX(90), c, 28, font=BODY, color="#B9C3BD", line_spacing=1.3)
    notes(sl, s)


def build_who(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    text(sl, PX(PADX), PX(230), PX(900), PX(40), s["eyebrow"], 26, font=UI, color=AMBER, caps=True, spacing=2)
    text(sl, PX(PADX), PX(280), PX(900), PX(160), s["h2"], 64, font=HEAD, bold=True, line_spacing=1.1)
    y = 480
    for b in s["bullets"]:
        rect(sl, PX(PADX), PX(y + 18), PX(14), PX(14), AMBER)
        text(sl, PX(PADX + 34), PX(y), PX(800), PX(50), b, 36, font=BODY, line_spacing=1.2)
        y += 68
    text(sl, PX(PADX), PX(y + 30), PX(820), PX(140), s["body"], 30, font=BODY, color="#B9C3BD", line_spacing=1.4)
    img = add_image(sl, s["image"], PX(1920 - 80 - 384), PX((1080 - 820) / 2), h=PX(820))
    if img is not None:
        img.left = PX(1920 - PADX) - img.width
    notes(sl, s)


def build_steps(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, PAPER)
    text(sl, PX(PADX), PX(PADT), PX(1200), PX(40), s["eyebrow"], 26, font=UI, color="#4A5551", caps=True, spacing=2)
    text(sl, PX(PADX), PX(146), PX(1500), PX(90), s["h2"], 64, font=HEAD, color=INK, bold=True, line_spacing=1.1)
    colw = (1920 - 2 * PADX - 3 * 40) / 4
    for i, (n, t, l, b) in enumerate(s["steps"]):
        x = PX(PADX + i * (colw + 40))
        hline(sl, x, PX(330), PX(colw), INK, 1.5)
        text(sl, x, PX(356), PX(colw), PX(40), n, 28, font=HEAD, color="#2274A5", bold=True)
        text(sl, x, PX(404), PX(colw), PX(56), t, 40, font=HEAD, color=INK, bold=True)
        text(sl, x, PX(468), PX(colw), PX(40), l, 24, font=UI, color="#4A5551", caps=True, spacing=1)
        text(sl, x, PX(520), PX(colw), PX(420), b, 28, font=BODY, color=INK, line_spacing=1.4)
    notes(sl, s)


def build_demo(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, PAPER)
    x0 = 80
    text(sl, PX(x0), PX(PADT), PX(520), PX(40), s["eyebrow"], 26, font=UI, color="#4A5551", caps=True, spacing=2)
    text(sl, PX(x0), PX(146), PX(520), PX(200), s["h2"], 54, font=HEAD, color=INK, bold=True, line_spacing=1.12)
    y = 380
    for p in s["paras"]:
        hline(sl, PX(x0), PX(y), PX(520), "#DDD7D4")
        text(sl, PX(x0), PX(y + 20), PX(520), PX(170), p, 28, font=BODY, color=INK, line_spacing=1.35)
        y += 200
    imgx, imgw = 660, 1920 - 80 - 660
    frame = rect(sl, PX(imgx) - PX(2), PX(PADT) - PX(2), PX(imgw) + PX(4), PX(860) + PX(4), "#FFFFFF", line="#DDD7D4", radius=PX(12))
    frame.shadow.inherit = True
    img = add_image(sl, s["image"], PX(imgx), PX(PADT), w=PX(imgw), fallback="trace-tasks-gate.png")
    if img is not None and img.height > PX(860):
        # keep the aspect ratio, cap the height, and shrink the frame to it
        ratio = PX(860) / img.height
        img.height = PX(860)
        img.width = int(img.width * ratio)
        frame.width = img.width + PX(4)
    elif img is not None:
        frame.height = img.height + PX(4)
    notes(sl, s)


NODE_W, NODE_H, NODE_GAP = 258, 96, 44


def node_card(sl, x, y, name: str, label: str, kind: str):
    color = {"code": "#B9C3BD", "agent": BLUE, "verify": GREEN, "gate": AMBER}[kind]
    w, h = NODE_W, NODE_H
    rect(sl, PX(x), PX(y), PX(w), PX(h), "#131F18", line="#2C3B33", radius=PX(12))
    rect(sl, PX(x), PX(y + 6), PX(6), PX(h - 12), color)
    # Node ids never wrap (underscores give no break opportunity): long ids get a smaller size instead of overflowing.
    name_pt = 26 if len(name) <= 11 else 23 if len(name) <= 14 else 20
    text(sl, PX(x + 18), PX(y + 14), PX(w - 24), PX(36), name, name_pt, font=HEAD, bold=True)
    text(sl, PX(x + 18), PX(y + 54), PX(w - 24), PX(30), label, 20 if len(label) <= 13 else 17, font=UI, color=color, caps=True, spacing=1)


def dashed(sl, x1, y1, x2, y2, color: str, weight: float = 2.0):
    ln = sl.shapes.add_connector(1, PX(x1), PX(y1), PX(x2), PX(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(weight)
    ln.line.dash_style = 4  # dash
    return ln


def arrow(sl, x, y):
    ln = sl.shapes.add_connector(1, PX(x), PX(y), PX(x + 44), PX(y))
    ln.line.color.rgb = rgb(CONNECTOR)
    ln.line.width = Pt(2)
    head = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, PX(x + 34), PX(y - 7), PX(11), PX(14))
    head.rotation = 90
    head.fill.solid()
    head.fill.fore_color.rgb = rgb(CONNECTOR)
    head.line.fill.background()
    head.shadow.inherit = False


def build_graph(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    text(sl, PX(PADX), PX(PADT), PX(1200), PX(40), s["eyebrow"], 26, font=UI, color=AMBER, caps=True, spacing=2)
    text(sl, PX(PADX), PX(146), PX(1600), PX(90), s["h2"], 64, font=HEAD, bold=True, line_spacing=1.1)

    def row(cols: list[list[tuple[str, str, str]]], y_mid: float, x: float = PADX) -> float:
        """Columns of one or two cards, centred on y_mid, joined by arrows at y_mid."""
        for ci, col in enumerate(cols):
            n = len(col)
            total_h = n * NODE_H + (n - 1) * 14
            y0 = y_mid - total_h / 2
            for j, (name, label, kind) in enumerate(col):
                node_card(sl, x, y0 + j * (NODE_H + 14), name, label, kind)
            if ci < len(cols) - 1:
                arrow(sl, x + NODE_W, y_mid)
            x += NODE_W + NODE_GAP
        return x

    y1 = 420
    end1 = row(s["row1"], y1)
    # repair arc from severity_check back over triage: a dashed bracket above the two cards
    step = NODE_W + NODE_GAP
    tri_x = PADX + 4 * step + NODE_W / 2
    sev_x = PADX + 5 * step + NODE_W / 2
    top = y1 - NODE_H / 2
    dashed(sl, sev_x, top, sev_x, top - 36, RED)
    dashed(sl, sev_x, top - 36, tri_x, top - 36, RED)
    dashed(sl, tri_x, top - 36, tri_x, top - 4, RED)
    head = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, PX(tri_x - 7), PX(top - 14), PX(14), PX(12))
    head.rotation = 180
    head.fill.solid()
    head.fill.fore_color.rgb = rgb(RED)
    head.line.fill.background()
    head.shadow.inherit = False
    text(sl, PX(tri_x - 150), PX(top - 76), PX(sev_x - tri_x + 300), PX(32), "repair ×1", 24, font=UI, color=RED, caps=True, align=PP_ALIGN.CENTER, spacing=1)
    y2 = 720
    end2 = row(s["row2"], y2)
    text(sl, PX(end2 + 12), PX(y2 - 100), PX(1920 - PADX - end2 - 12), PX(240), s["explainer"], 26, font=BODY, color="#B9C3BD", line_spacing=1.4)
    lx = PADX
    for name, color in s["legend"]:
        rect(sl, PX(lx), PX(966), PX(18), PX(18), color)
        text(sl, PX(lx + 30), PX(960), PX(260), PX(32), name, 24, font=UI, color="#B9C3BD", caps=True, spacing=1)
        lx += 40 + 30 + len(name) * 14 + 40
    notes(sl, s)


def build_cards(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    text(sl, PX(PADX), PX(PADT), PX(1200), PX(40), s["eyebrow"], 26, font=UI, color=AMBER, caps=True, spacing=2)
    text(sl, PX(PADX), PX(146), PX(1600), PX(90), s["h2"], 64, font=HEAD, bold=True, line_spacing=1.1)
    cw = (1920 - 2 * PADX - 2 * 24) / 3
    ch = 320
    for i, (t, b) in enumerate(s["cards"]):
        col, rw = i % 3, i // 3
        x = PADX + col * (cw + 24)
        y = 280 + rw * (ch + 24)
        rect(sl, PX(x), PX(y), PX(cw), PX(ch), "#131F18", line="#2C3B33", radius=PX(16))
        text(sl, PX(x + 32), PX(y + 28), PX(cw - 64), PX(34), t, 24, font=UI, color=AMBER, caps=True, spacing=1)
        text(sl, PX(x + 32), PX(y + 76), PX(cw - 64), PX(ch - 90), b, 26, font=BODY, line_spacing=1.35)
    text(sl, PX(PADX), PX(972), PX(1700), PX(70), s["footnote"], 22, font=BODY, color="#B9C3BD", line_spacing=1.35)
    notes(sl, s)


def build_numbers(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    text(sl, PX(PADX), PX(PADT), PX(1400), PX(40), s["eyebrow"], 26, font=UI, color=AMBER, caps=True, spacing=2)
    text(sl, PX(PADX), PX(146), PX(1600), PX(90), s["h2"], 64, font=HEAD, bold=True, line_spacing=1.1)
    cw = (1920 - 2 * PADX - 2 * 48) / 3
    for i, (n, c) in enumerate(s["numbers"]):
        col, rw = i % 3, i // 3
        x = PADX + col * (cw + 48)
        y = 300 + rw * 300
        hline(sl, PX(x), PX(y), PX(cw), "#2C3B33")
        text(sl, PX(x), PX(y + 28), PX(cw), PX(110), n, 104, font=HEAD, bold=True, line_spacing=1.0)
        text(sl, PX(x), PX(y + 150), PX(cw), PX(90), c, 28, font=BODY, color="#B9C3BD", line_spacing=1.3)
    text(sl, PX(PADX), PX(930), PX(1500), PX(100), s["footer"], 26, font=BODY, color="#B9C3BD", line_spacing=1.4)
    notes(sl, s)


def build_closing(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    logo(sl, PX(960 - 175), PX(150), 96, 31, 6, word_px=56)
    text(sl, PX(240), PX(330), PX(1440), PX(260), s["quote"], 60, font=BODY, italic=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
    text(sl, PX(PADX), PX(680), PX(1700), PX(50), s["mono"], 32, font=MONO, color=BLUE, align=PP_ALIGN.CENTER)
    text(sl, PX(PADX), PX(740), PX(1700), PX(50), s["mono2"], 32, font=MONO, color=BLUE, align=PP_ALIGN.CENTER)
    text(sl, PX(PADX), PX(820), PX(1700), PX(40), s["eyebrow"], 26, font=UI, color=AMBER, caps=True, align=PP_ALIGN.CENTER, spacing=2)
    notes(sl, s)


BUILDERS = {"title": build_title, "problem": build_problem, "who": build_who, "steps": build_steps, "demo": build_demo, "graph": build_graph, "cards": build_cards, "numbers": build_numbers, "closing": build_closing}


def build() -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    for s in SLIDES:
        BUILDERS[s["kind"]](prs, s)
    os.makedirs(OUT_DIR, exist_ok=True)
    prs.save(OUT)
    print(f"deck: {OUT} ({len(SLIDES)} slides)")
    return OUT


def export(pptx_path: str) -> str:
    """Render every slide to PNG through PowerPoint (COM automation; Windows with Office)."""
    import win32com.client  # type: ignore

    out_dir = os.path.join(OUT_DIR, "preview")
    os.makedirs(out_dir, exist_ok=True)
    for f in os.listdir(out_dir):
        os.remove(os.path.join(out_dir, f))
    # Early binding (generated typelib wrapper): the dynamic dispatch cannot set Presentation.EmbedTrueTypeFonts.
    app = win32com.client.gencache.EnsureDispatch("PowerPoint.Application")
    pres = app.Presentations.Open(os.path.abspath(pptx_path), False, False, False)
    try:
        # Embed the brand fonts (OFL, embeddable) so the deck renders the same on a machine without them installed.
        try:
            pres.EmbedTrueTypeFonts = True
            pres.Save()
            print("fonts embedded")
        except Exception as e:  # noqa: BLE001
            print(f"font embedding skipped: {e}")
        pres.Export(os.path.abspath(out_dir), "PNG", 1920, 1080)
    finally:
        pres.Close()
        app.Quit()
    print(f"preview: {out_dir}")
    return out_dir


if __name__ == "__main__":
    path = build()
    if "--export" in sys.argv:
        export(path)
