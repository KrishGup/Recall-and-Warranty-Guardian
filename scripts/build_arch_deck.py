"""Build the Guardian architecture deck as PowerPoint from scripts/arch_deck_content.py.

    python scripts/build_arch_deck.py                 # -> docs/deck/Guardian-Architecture.pptx
    python scripts/build_arch_deck.py --export        # also renders every slide to docs/deck/arch-preview/NN.png (PowerPoint, Windows)

Same canvas, fonts, and colors as the demo deck (build_deck.py): 16:9 at 13.333 x 7.5 in, a 1920x1080 design grid,
Roboto Slab, Lora, Habibi, Consolas. Ink slides carry the diagrams; paper slides carry the tables and columns.
The slide kinds are: title, columns, table, overview, graph, cards, gate, deployment, closing."""
from __future__ import annotations

import os
import pathlib
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from build_deck import BODY, HEAD, MONO, PADT, PADX, UI, FS, PX, arrow, background, build_cards, build_closing, build_graph, build_title, dashed, hline, node_card, notes, rect, rgb, text  # noqa: E402
from deck_content import AMBER, BLUE, GREEN, INK, MUTED_DARK, PAPER, RED  # noqa: E402
from arch_deck_content import SLIDES  # noqa: E402

OUT_DIR = os.path.join(ROOT, "docs", "deck")
OUT = os.path.join(OUT_DIR, "Guardian-Architecture.pptx")
PREVIEW = os.path.join(OUT_DIR, "arch-preview")

MUTED = "#4A5551"
RULE = "#DDD7D4"
TINT = "#EFEBE9"
STEP_BLUE = "#2274A5"
CARD_DARK = "#131F18"
RULE_DARK = "#2C3B33"
WIDTH = 1920 - 2 * PADX


def footer(sl, n: int, total: int, dark: bool):
    text(sl, PX(PADX), PX(1018), PX(WIDTH), PX(30), f"Guardian · Architecture · {n} / {total}", 20, font=UI, color=MUTED_DARK if dark else MUTED, align=PP_ALIGN.RIGHT, caps=True, spacing=1)


def heading(sl, s: dict, dark: bool):
    text(sl, PX(PADX), PX(PADT), PX(1400), PX(40), s["eyebrow"], 26, font=UI, color=AMBER if dark else MUTED, caps=True, spacing=2)
    text(sl, PX(PADX), PX(146), PX(WIDTH), PX(90), s["h2"], 60, font=HEAD, color=PAPER if dark else INK, bold=True, line_spacing=1.08)


_LORA = None


def _lora(size: float):
    """The Lora face at `size` px, from design/fonts, for exact advance widths. None when the file is missing."""
    global _LORA
    if _LORA is None:
        try:
            from PIL import ImageFont
            path = pathlib.Path(ROOT, "design", "fonts", "Lora-Regular.ttf")
            _LORA = (ImageFont, str(path)) if path.exists() else False
        except ImportError:
            _LORA = False
    if not _LORA:
        return None
    return _LORA[0].truetype(_LORA[1], max(1, round(size)))


def est_width(t: str, size: float) -> float:
    """Rendered width of `t` in Lora at `size` px. Exact from the font file (within 2 % of the PowerPoint export),
    with a 3 % safety factor; a per-glyph-class estimate when PIL or the font is missing."""
    face = _lora(size)
    if face is not None:
        return face.getlength(t) * 1.03
    w = 0.0
    for ch in t:
        if ch == " ":
            f = 0.27
        elif ch.isupper():
            f = 0.70
        elif ch.isdigit():
            f = 0.59
        elif ch in "iljtfr.,;:'!|":
            f = 0.33
        elif ch in "mw":
            f = 0.88
        elif ch in "-()[]/\\":
            f = 0.38
        elif ch == "_":
            f = 0.55
        elif ch.islower():
            f = 0.55
        else:
            f = 0.60
        w += f * size
    return w


def est_lines(t: str, w: float, size: float) -> int:
    """Greedy word wrap on the estimated widths: how many lines `t` takes in a box `w` px wide. Long tokens such as
    GUARDIAN_DAILY_BUDGET_USD do not break, so a per-character average undercounts them; this counts them as words."""
    lines, cur = 1, 0.0
    space = est_width(" ", size)
    for word in t.split(" "):
        ww = est_width(word, size)
        if cur == 0:
            cur = ww
        elif cur + space + ww <= w:
            cur += space + ww
        else:
            lines += 1
            cur = ww
    return lines


LORA_LINE = 1.2  # the natural line height of Lora, in em; PowerPoint multiplies it by the paragraph line spacing


def bullets(sl, x: float, y: float, w: float, items: list[str], size: float, color: str, gap: float, marker: str, line_spacing: float):
    """Bullet rows with a small square marker; returns the y after the last row. Row height comes from a word-wrap
    estimate at the real line pitch, so the following row never overlaps a wrapped one."""
    for t in items:
        lines_n = est_lines(t, (w - 30) * 0.97, size)
        h = lines_n * size * LORA_LINE * line_spacing
        rect(sl, PX(x), PX(y + size * 0.55), PX(12), PX(12), marker)
        text(sl, PX(x + 30), PX(y), PX(w - 30), PX(h), t, size, font=BODY, color=color, line_spacing=line_spacing)
        y += h + gap
    return y


# ------------------------------------------------------------------------------------------------ builders
def build_columns(prs, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, PAPER)
    heading(sl, s, False)
    cols = s["columns"]
    gap = 48
    colw = (WIDTH - gap * (len(cols) - 1)) / len(cols)
    body_size = 26 if len(cols) == 2 else 23
    top = 300
    bottom = 980
    if s.get("stats"):
        bottom = 800
    for i, (label, items) in enumerate(cols):
        x = PADX + i * (colw + gap)
        hline(sl, PX(x), PX(top), PX(colw), INK, 1.5)
        text(sl, PX(x), PX(top + 22), PX(colw), PX(36), label, 24, font=UI, color=STEP_BLUE, caps=True, spacing=2)
        bullets(sl, x, top + 76, colw, items, body_size, INK, 18, STEP_BLUE, 1.3)
    if s.get("stats"):
        hline(sl, PX(PADX), PX(bottom + 20), PX(WIDTH), RULE)
        sw = (WIDTH - 2 * 48) / 3
        for i, (num, cap) in enumerate(s["stats"]):
            x = PADX + i * (sw + 48)
            text(sl, PX(x), PX(bottom + 46), PX(sw), PX(90), num, 80, font=HEAD, color=INK, bold=True, line_spacing=1.0)
            text(sl, PX(x), PX(bottom + 140), PX(sw), PX(60), cap, 24, font=BODY, color=MUTED, line_spacing=1.25)
    footer(sl, n, total, False)
    notes(sl, s)


def build_table(prs, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, PAPER)
    heading(sl, s, False)
    header, rows, widths = s["header"], s["rows"], s["widths"]
    assert abs(sum(widths) - WIDTH) < 2, f"{s['id']}: column widths sum to {sum(widths)}, want {WIDTH}"
    nrows = len(rows) + 1
    head_h = 46
    avail = 1000 - 280

    def heights(size: float) -> list[float]:
        """Row heights from the longest cell of each row (characters per line from the column width)."""
        out = []
        for r in rows:
            lines_n = 1
            for w, cell in zip(widths, r):
                cpl = max(8, int((w - 28) / (size * 0.52)))
                lines_n = max(lines_n, -(-len(cell) // cpl))
            out.append(lines_n * size * 1.3 + 16)
        return out

    # The largest body size whose estimated table fits above the footer; the export is the final check.
    for body_size in (25, 24, 23, 22, 21, 20):
        row_h = heights(body_size)
        total_h = head_h + sum(row_h)
        if total_h <= avail:
            break
    else:
        raise SystemExit(f"{s['id']}: table needs {total_h:.0f}px at 20px, {avail}px available. Shorten the rows or split the slide.")
    shape = sl.shapes.add_table(nrows, len(widths), PX(PADX), PX(280), PX(WIDTH), PX(total_h))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for ci, w in enumerate(widths):
        tbl.columns[ci].width = PX(w)
    tbl.rows[0].height = PX(head_h)
    for ri, h in enumerate(row_h, start=1):
        tbl.rows[ri].height = PX(h)

    def cell_text(cell, value: str, size: float, font: str, color: str, bold=False, caps=False, fill: str | None = None):
        cell.margin_left = cell.margin_right = PX(14)
        cell.margin_top = PX(8)
        cell.margin_bottom = PX(6)
        cell.vertical_anchor = MSO_ANCHOR.TOP
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = value.upper() if caps else value
        run.font.name = font
        run.font.size = FS(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if fill:
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)

    for ci, h in enumerate(header):
        cell_text(tbl.cell(0, ci), h, 21, UI, PAPER, caps=True, fill=INK)
    for ri, r in enumerate(rows, start=1):
        fill = "#FFFFFF" if ri % 2 else TINT
        for ci, value in enumerate(r):
            first = ci == 0 and len(widths) > 2
            cell_text(tbl.cell(ri, ci), value, body_size, HEAD if first else BODY, INK, bold=first, fill=fill)
    footer(sl, n, total, False)
    notes(sl, s)


def ink_box(sl, x, y, w, h, title: str, body: str, border: str = RULE_DARK, title_color: str = PAPER, body_size: float = 21, title_size: float = 26):
    rect(sl, PX(x), PX(y), PX(w), PX(h), CARD_DARK, line=border, radius=PX(12), line_w=1.25)
    text(sl, PX(x + 22), PX(y + 18), PX(w - 44), PX(36), title, title_size, font=HEAD, color=title_color, bold=True)
    text(sl, PX(x + 22), PX(y + 58), PX(w - 44), PX(h - 70), body, body_size, font=BODY, color=MUTED_DARK, line_spacing=1.3)


def label(sl, x, y, w, t: str, color: str = MUTED_DARK, align=PP_ALIGN.LEFT):
    text(sl, PX(x), PX(y), PX(w), PX(30), t, 22, font=UI, color=color, caps=True, spacing=1, align=align)


def varrow(sl, x, y1, y2, color: str = "#4A5551", dash: bool = False):
    ln = sl.shapes.add_connector(1, PX(x), PX(y1), PX(x), PX(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(2)
    if dash:
        ln.line.dash_style = 4
    down = y2 > y1
    head = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, PX(x - 7), PX((y2 - 12) if down else y2), PX(14), PX(12))
    head.rotation = 0 if not down else 180
    head.fill.solid()
    head.fill.fore_color.rgb = rgb(color)
    head.line.fill.background()
    head.shadow.inherit = False


def harrow(sl, x1, x2, y, color: str = "#4A5551", dash: bool = False):
    ln = sl.shapes.add_connector(1, PX(x1), PX(y), PX(x2 - 12), PX(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(2)
    if dash:
        ln.line.dash_style = 4
    head = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, PX(x2 - 12), PX(y - 7), PX(12), PX(14))
    head.rotation = 90
    head.fill.solid()
    head.fill.fore_color.rgb = rgb(color)
    head.line.fill.background()
    head.shadow.inherit = False


def build_overview(prs, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    heading(sl, s, True)
    # Row 1: inputs -> the engine -> the household
    y1, h1 = 290, 200
    ink_box(sl, PADX, y1, 380, h1, "Inputs", "Recall feeds: CPSC, NHTSA by vehicle, openFDA. Receipts, order emails, and VINs from the household.", body_size=20)
    harrow(sl, PADX + 380, PADX + 440, y1 + h1 / 2)
    gx, gw = PADX + 440, 820
    rect(sl, PX(gx), PX(y1), PX(gw), PX(h1), CARD_DARK, line=BLUE, radius=PX(14), line_w=1.5)
    text(sl, PX(gx + 22), PX(y1 + 18), PX(gw - 44), PX(34), "gren on Strands Agents", 26, font=HEAD, color=PAPER, bold=True)
    label(sl, gx + 22, y1 + 56, gw - 44, "two graphs · budgets · retries · gates as interrupts", BLUE)
    node_card(sl, gx + 22, y1 + 92, "nightly-sweep", "13 nodes · gate", "agent")
    node_card(sl, gx + 22 + 258 + 24, y1 + 92, "intake", "3 nodes", "agent")
    text(sl, PX(gx + 22 + 2 * 282 + 10), PX(y1 + 92), PX(gw - 44 - 2 * 282 - 10), PX(96), "Typed output on every agent node. Artifacts for every run.", 18, font=BODY, color=MUTED_DARK, line_spacing=1.3)
    harrow(sl, gx + gw, gx + gw + 60, y1 + h1 / 2, AMBER, dash=True)
    hx = gx + gw + 60
    ink_box(sl, hx, y1, 1920 - PADX - hx, h1, "Household", "One SMS with one decision, or the dashboard card. The answer resumes the run.", border=AMBER, title_color=AMBER, body_size=20)
    # Row 2: store -> API -> dashboard
    y2, h2 = 570, 190
    varrow(sl, gx + gw / 2, y1 + h1 + 4, y2 - 4)
    label(sl, gx + gw / 2 + 14, y1 + h1 + 30, 400, "events · state · artifacts")
    ink_box(sl, PADX, y2, 520, h2, "Household store", "items, recalls, matches, decisions, activity, sweeps, outbox. JSON files shaped like the DynamoDB plan. Run store: state, events, artifacts.", body_size=20)
    harrow(sl, PADX + 520, PADX + 580, y2 + h2 / 2)
    ink_box(sl, PADX + 580, y2, 520, h2, "Guardian API", "FastAPI and SSE. Turns gates into decisions, records answers, resumes the run. Mounts the gren run API at /gren. Token-gated writes.", body_size=20)
    harrow(sl, PADX + 1100, PADX + 1160, y2 + h2 / 2)
    ink_box(sl, PADX + 1160, y2, WIDTH - 1160, h2, "Dashboard", "Vite and React. Home, Decisions, Inventory, Activity, Settings, Agent flow: the trace workbench with graph, timeline, events, gate, inspector, fork.", body_size=20)
    varrow(sl, PADX + 1160 + (WIDTH - 1160) / 2, y2 - 4, y1 + h1 + 4, AMBER, dash=True)
    # Row 3: AWS band
    y3, h3 = 800, 176
    rect(sl, PX(PADX), PX(y3), PX(WIDTH), PX(h3), CARD_DARK, line=RULE_DARK, radius=PX(14), line_w=1.25)
    label(sl, PADX + 22, y3 + 14, 200, "AWS", AMBER)
    cells = [
        ("Amazon Bedrock", "Claude Haiku 4.5, Sonnet 5, Opus 5 through global inference profiles."),
        ("Amazon S3", "The state bucket: both stores mirror after each change. Releases and config."),
        ("Amazon SES", "The remedy email with the receipt attached. Outbox until the sender is verified."),
        ("EC2 + Caddy", "One instance, automatic HTTPS, a systemd timer at 06:00 UTC. Live."),
        ("Bedrock AgentCore", "Runtime entrypoint built and tested locally. Next: needs the account quotas."),
    ]
    cw = (WIDTH - 44 - 4 * 24) / 5
    for i, (t, b) in enumerate(cells):
        x = PADX + 22 + i * (cw + 24)
        text(sl, PX(x), PX(y3 + 50), PX(cw), PX(34), t, 24, font=HEAD, color=PAPER, bold=True)
        text(sl, PX(x), PX(y3 + 90), PX(cw), PX(80), b, 19, font=BODY, color=MUTED_DARK, line_spacing=1.3)
    for x in (PADX + 260, PADX + 840, PADX + 1380):
        dashed(sl, x, y2 + h2, x, y3, "#4A5551", 1.5)
    footer(sl, n, total, True)
    notes(sl, s)


def build_gate(prs, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    heading(sl, s, True)
    steps = s["steps"]
    gap = 40
    w = (WIDTH - gap * (len(steps) - 1)) / len(steps)
    y, ch = 290, 300
    for i, (t, b) in enumerate(steps):
        x = PADX + i * (w + gap)
        colr = AMBER if i in (0, 3) else BLUE if i == 4 else PAPER
        rect(sl, PX(x), PX(y), PX(w), PX(ch), CARD_DARK, line=AMBER if i in (0, 3) else RULE_DARK, radius=PX(12), line_w=1.25)
        text(sl, PX(x + 20), PX(y + 18), PX(w - 40), PX(34), f"{i + 1}", 26, font=HEAD, color=STEP_BLUE if i not in (0, 3) else AMBER, bold=True)
        text(sl, PX(x + 20), PX(y + 56), PX(w - 40), PX(36), t, 26, font=HEAD, color=colr, bold=True)
        text(sl, PX(x + 20), PX(y + 100), PX(w - 40), PX(ch - 116), b, 20, font=BODY, color=MUTED_DARK, line_spacing=1.3)
        if i < len(steps) - 1:
            harrow(sl, x + w + 4, x + w + gap - 4, y + ch / 2, AMBER if i in (0, 2, 3) else "#4A5551", dash=i in (0, 2))
    label(sl, PADX, 630, 600, "Rules the engine enforces", AMBER)
    rules = s["rules"]
    rw = (WIDTH - 24 * (len(rules) - 1)) / len(rules)
    for i, (t, b) in enumerate(rules):
        x = PADX + i * (rw + 24)
        ink_box(sl, x, 670, rw, 270, t, b, body_size=20, title_size=24)
    footer(sl, n, total, True)
    notes(sl, s)


def build_deployment(prs, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    background(sl, INK)
    heading(sl, s, True)
    y, h = 290, 240
    ink_box(sl, PADX, y, 380, h, "deploy_ec2.py", "up, update, status, logs, run, down. Builds the dashboard, uploads the release and config/env, creates the role, the security group, the elastic IP, and the instance.", body_size=19)
    harrow(sl, PADX + 380, PADX + 440, y + h / 2)
    ink_box(sl, PADX + 440, y, 380, h, "S3 state bucket", "releases/latest.zip and config/env for the instance. guardian/household and guardian/runs: the mirrored stores.", body_size=19)
    harrow(sl, PADX + 820, PADX + 880, y + h / 2)
    ix, iw, ih = PADX + 880, WIDTH - 880, 470
    rect(sl, PX(ix), PX(y - 10), PX(iw), PX(ih), CARD_DARK, line=AMBER, radius=PX(14), line_w=1.5)
    text(sl, PX(ix + 22), PX(y + 8), PX(iw - 44), PX(34), "EC2 t3.small · Amazon Linux 2023", 26, font=HEAD, color=PAPER, bold=True)
    label(sl, ix + 22, y + 46, iw - 44, "install.sh runs on the first boot and on every update", AMBER)
    inner = [
        ("Caddy :443", "Automatic Let's Encrypt certificate for guardian.<ip>.sslip.io. Reverse proxy to 8787."),
        ("guardian.service :8787", "Python 3.12 venv with gren and guardian. FastAPI, SSE, the built dashboard, state sync on."),
        ("guardian-sweep.timer", "POST /api/sweep with the token at 06:00 UTC (02:00 New York). The service runs with --no-schedule."),
        ("IAM role guardian-ec2", "Bedrock invoke, the state bucket, SES send, SSM. No SSH: logs and commands run through SSM."),
    ]
    cw, chh = (iw - 44 - 24) / 2, 170
    for i, (t, b) in enumerate(inner):
        cx = ix + 22 + (i % 2) * (cw + 24)
        cy = y + 90 + (i // 2) * (chh + 20)
        rect(sl, PX(cx), PX(cy), PX(cw), PX(chh), INK, line=RULE_DARK, radius=PX(10))
        text(sl, PX(cx + 18), PX(cy + 14), PX(cw - 36), PX(32), t, 22, font=MONO, color=BLUE, bold=True)
        text(sl, PX(cx + 18), PX(cy + 52), PX(cw - 36), PX(chh - 62), b, 19, font=BODY, color=MUTED_DARK, line_spacing=1.3)
    y2, h2 = 800, 176
    ink_box(sl, PADX, y2, 380, h2, "Household browser", "guardian.44-214-230-44.sslip.io. Reads are open. Writes need the token, set one time by ?token=.", border=AMBER, title_color=AMBER, body_size=19)
    ink_box(sl, PADX + 440, y2, 380, h2, "Amazon Bedrock", "Every agent node calls a Claude inference profile with the instance role. Two caps bound the spend.", border=BLUE, body_size=19)
    ink_box(sl, PADX + 880, y2, WIDTH - 880, h2, "Bedrock AgentCore · next", "deploy/agentcore: a BedrockAgentCoreApp entrypoint, CodeZip runtime, S3 state around every invocation. agentcore deploy fails with maxAgents limit exceeded until AWS Support raises the quotas from 0.", border=RULE_DARK, title_color=MUTED_DARK, body_size=19)
    varrow(sl, PADX + 190, y2 - 4, y + h + 4, AMBER, dash=True)
    label(sl, PADX + 204, y + h + 30, 220, "https", AMBER)
    varrow(sl, PADX + 630, y + h + 4, y2 - 4, BLUE, dash=True)
    label(sl, PADX + 644, y + h + 30, 300, "model calls", BLUE)
    footer(sl, n, total, True)
    notes(sl, s)


def build_title_slide(prs, s, n, total):
    build_title(prs, s)


def build_graph_slide(prs, s, n, total):
    build_graph(prs, s)
    footer(prs.slides[-1], n, total, True)


def build_cards_slide(prs, s, n, total):
    build_cards(prs, s)
    footer(prs.slides[-1], n, total, True)


def build_closing_slide(prs, s, n, total):
    build_closing(prs, s)


BUILDERS = {
    "title": build_title_slide, "columns": build_columns, "table": build_table, "overview": build_overview, "graph": build_graph_slide,
    "cards": build_cards_slide, "gate": build_gate, "deployment": build_deployment, "closing": build_closing_slide,
}


def build() -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = PX(1920), PX(1080)
    total = len(SLIDES)
    for i, s in enumerate(SLIDES, start=1):
        BUILDERS[s["kind"]](prs, s, i, total)
    os.makedirs(OUT_DIR, exist_ok=True)
    prs.save(OUT)
    print(f"deck: {OUT} ({total} slides)")
    return OUT


def export(pptx_path: str, pdf: bool = False) -> str:
    """Open the deck in PowerPoint, embed the brand fonts, write one PNG per slide to docs/deck/arch-preview/, and with
    `pdf` also save docs/deck/Guardian-Architecture.pdf (the copy for a reader without PowerPoint or the fonts)."""
    import win32com.client  # type: ignore

    os.makedirs(PREVIEW, exist_ok=True)
    for f in os.listdir(PREVIEW):
        os.remove(os.path.join(PREVIEW, f))
    app = win32com.client.gencache.EnsureDispatch("PowerPoint.Application")
    pres = app.Presentations.Open(os.path.abspath(pptx_path), False, False, False)
    try:
        # Save() ignores EmbedTrueTypeFonts on a file it has not changed; SaveAs with the embed flag writes ppt/fonts/*.
        try:
            pres.SaveAs(os.path.abspath(pptx_path), 24, -1)  # ppSaveAsOpenXMLPresentation, EmbedTrueTypeFonts=msoTrue
            print("fonts embedded")
        except Exception as e:  # noqa: BLE001
            print(f"font embedding skipped: {e}")
        pres.Export(os.path.abspath(PREVIEW), "PNG", 1920, 1080)
        if pdf:
            pdf_path = os.path.splitext(os.path.abspath(pptx_path))[0] + ".pdf"
            pres.SaveAs(pdf_path, 32)  # ppSaveAsPDF
            print(f"pdf: {pdf_path}")
    finally:
        pres.Close()
        app.Quit()
    print(f"preview: {PREVIEW}")
    return PREVIEW


if __name__ == "__main__":
    path = build()
    if "--export" in sys.argv or "--pdf" in sys.argv:
        export(path, pdf="--pdf" in sys.argv)
